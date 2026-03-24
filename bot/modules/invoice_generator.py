import os
import random
from datetime import datetime
from docx import Document
from docx2pdf import convert

TEMPLATE_PATH = "bot/templates/invoice_template.docx"
BASE_OUTPUT_DIR = "generated_docs/invoice"

def replace_placeholders(paragraphs, replacements):
    for paragraph in paragraphs:
        for key, value in replacements.items():
            if f"{{{{{key}}}}}" in paragraph.text:
                for run in paragraph.runs:
                    if f"{{{{{key}}}}}" in run.text:
                        run.text = run.text.replace(f"{{{{{key}}}}}", value)

def replace_in_tables(tables, replacements):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                replace_placeholders(cell.paragraphs, replacements)

def generate_invoice(data: list) -> tuple:
    if len(data) != 11:
        raise ValueError("Неверное количество строк. Ожидалось 11.")

    (
        contract_number, contract_date, recipient_inn, recipient_bank,
        bank_inn, bank_kpp, recipient_name, bik,
        corr_account, account_number, amount
    ) = data

    # Для шаблона — дата договора
    try:
        datetime.strptime(contract_date.strip(), "%d.%m.%Y")
    except ValueError:
        raise ValueError("Неверный формат даты. Используйте ДД.ММ.ГГГГ.")

    # Для имени файла и папки — сегодняшняя дата
    today = datetime.today()
    today_folder = today.strftime("%Y-%m-%d")
    today_str = today.strftime("%Y-%m-%d")
    unique_suffix = str(random.getrandbits(24))[-6:]

    filename_base = f"Invoice_{contract_number}_{today_str}_{unique_suffix}"
    output_dir = os.path.join(BASE_OUTPUT_DIR, today_folder)
    os.makedirs(output_dir, exist_ok=True)

    docx_path = os.path.join(output_dir, f"{filename_base}.docx")
    pdf_path = os.path.join(output_dir, f"{filename_base}.pdf")

    replacements = {
        "ContractNumber": contract_number,
        "ContractDate": contract_date,
        "RecipientINN": recipient_inn,
        "RecipientBank": recipient_bank,
        "BankINN": bank_inn,
        "BankKPP": bank_kpp,
        "RecipientName": recipient_name,
        "BIK": bik,
        "CorrAccount": corr_account,
        "AccountNumber": account_number,
        "Amount": amount
    }

    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Шаблон не найден: {TEMPLATE_PATH}")

    doc = Document(TEMPLATE_PATH)
    replace_placeholders(doc.paragraphs, replacements)
    replace_in_tables(doc.tables, replacements)
    doc.save(docx_path)

    try:
        convert(docx_path, pdf_path)
    except Exception as e:
        print(f"[Ошибка при конвертации PDF]: {e}")
        pdf_path = None

    return docx_path, pdf_path
